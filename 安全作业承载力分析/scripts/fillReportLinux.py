# -*- coding: utf-8 -*-
"""Linux 报告填充：无 Word/WPS COM 时，用 python-docx 套模板并插入 PNG。

与 fillReport.py + pngCharts.py apply 对齐的能力：
  - replace：段落与单元格全文子串替换（长键优先）
  - paragraph_replace：anchor + occurrence，可选 next_paragraph 整段覆盖
  - tables：按列名或 anchor 定位，覆盖表头、删示例行、写入数据行
  - 矩阵表 max_units_per_table 纵切，块间插入「（续表）」
  - 空明细表删除
  - trim_month_weeks
  - 日/周：替换模板承载力图位，并在引导句/评述句后插入作业类型图与管理承载力图
  - 月：按「图 N…」图题把 PNG 写入其前一段

不改计算公式。正文段落的首个 run 格式（字体、字号）保留。
"""
from __future__ import annotations

import argparse
import logging
import os
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from typing import Any, Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from capacity.office_chart import (
    DETAIL_FONT_PT, DETAIL_LINE_PT, TEXT_WIDTH_FALLBACK_PT, detail_col_widths)
from fillReport import TEMPLATES

LOGGER = logging.getLogger("fillReportLinux")

_BROKEN_REF = "Error: Reference source not found"
_WEEK_RE = re.compile(r"^([1-5])\..*第([1-5])周")
_SINGLE_PREFIX = {"day": "cap_d1", "week": "cap_w1"}
_WEEK_PREFIX = (("第1周", "w1"), ("第2周", "w2"), ("第3周", "w3"),
                ("第4周", "w4"), ("第5周", "w5"))
_MANAGE_ANCHOR = {"day": "本日管理承载力", "week": "本周管理承载力",
                  "month": "全月管理承载力"}
_PTYPE_ANCHOR = "各作业类型日计划数量如下图所示"
_CHART_W_PT = 435.0
_MAX_H_PT = 240.0
_MATRIX_FONT_PT = 12.0
_MATRIX_FIRST_PT = 60.0


def _substitute(text: str, mapping: Dict[str, str]) -> str:
    """长键优先的纯子串替换。"""
    if not text or not mapping:
        return text
    for key in sorted(mapping, key=len, reverse=True):
        if key and key in text:
            text = text.replace(key, mapping[key])
    return text


def _run_has_drawing(run_elm) -> bool:
    return (run_elm.find(qn("w:drawing")) is not None
            or run_elm.find(qn("w:pict")) is not None)


def set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """改段内文字，保留首个文本 run 的 rPr 与段落 pPr。含图的 run 不动。"""
    text_runs = [r for r in paragraph.runs if not _run_has_drawing(r._element)]
    if not text_runs:
        paragraph.add_run(new_text)
        return
    text_runs[0].text = new_text
    for run in text_runs[1:]:
        parent = run._element.getparent()
        if parent is not None:
            parent.remove(run._element)


def _iter_paragraphs(doc: Document) -> List[Paragraph]:
    return list(doc.paragraphs)


def _cell_paragraphs(cell) -> List[Paragraph]:
    return list(cell.paragraphs)


def apply_replace(doc: Document, mapping: Dict[str, str]) -> None:
    """全文子串替换：正文段落 + 表格单元格段落。"""
    if not mapping:
        return
    for paragraph in _iter_paragraphs(doc):
        old = paragraph.text
        new = _substitute(old, mapping)
        if new != old:
            set_paragraph_text(paragraph, new)
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                cid = id(cell._tc)
                if cid in seen:
                    continue
                seen.add(cid)
                for paragraph in _cell_paragraphs(cell):
                    old = paragraph.text
                    new = _substitute(old, mapping)
                    if new != old:
                        set_paragraph_text(paragraph, new)


def _paragraph_index(doc: Document) -> List[Paragraph]:
    return list(doc.paragraphs)


def apply_paragraph_replace(doc: Document, specs: Sequence[Dict[str, Any]]) -> None:
    """anchor 第 N 次命中：子串替换，或 next_paragraph 时覆盖下一段正文。"""
    if not specs:
        return
    groups: Dict[str, Dict[int, Dict[str, Any]]] = {}
    for spec in specs:
        anchor = spec.get("anchor")
        if not anchor or spec.get("new_text") is None:
            continue
        groups.setdefault(str(anchor), {})[int(spec.get("occurrence", 1))] = spec

    for anchor, occ_map in groups.items():
        seen = 0
        paragraphs = _paragraph_index(doc)
        for index, paragraph in enumerate(paragraphs):
            text = paragraph.text
            start = 0
            while True:
                pos = text.find(anchor, start)
                if pos < 0:
                    break
                seen += 1
                spec = occ_map.get(seen)
                if spec is not None:
                    if spec.get("next_paragraph"):
                        if index + 1 < len(paragraphs):
                            set_paragraph_text(paragraphs[index + 1],
                                               str(spec["new_text"]))
                    else:
                        text = text[:pos] + str(spec["new_text"]) + text[pos + len(anchor):]
                        set_paragraph_text(paragraph, text)
                start = pos + max(len(str(spec["new_text"])) if spec and not spec.get("next_paragraph") else len(anchor), 1)
                if spec and not spec.get("next_paragraph"):
                    # 文本已重写，按新串继续，避免死循环
                    text = paragraph.text
                    start = pos + len(str(spec["new_text"]))


def repair_broken_refs(doc: Document) -> None:
    """LibreOffice 把模板交叉引用收成 Error 文本时，月总评述收成「图 2」。"""
    for paragraph in _iter_paragraphs(doc):
        if _BROKEN_REF in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(_BROKEN_REF, "图 2"))


def _header_texts(table: Table) -> List[str]:
    if not table.rows:
        return []
    seen = set()
    heads = []
    for cell in table.rows[0].cells:
        cid = id(cell._tc)
        if cid in seen:
            continue
        seen.add(cid)
        heads.append(cell.text.replace("\n", "").strip())
    return heads


def _match_score(heads: Sequence[str], columns: Sequence[str]) -> int:
    cset = set(columns)
    score = 0
    for head in heads[:len(columns)]:
        if head and head in cset:
            score += 2
    run = 0
    prev = -10
    for index, head in enumerate(heads):
        if head in cset:
            if index == prev + 1:
                run += 1
                score += run
            else:
                run = 1
            prev = index
    if len(heads) == len(columns):
        score += 20
    return score


def _body_children(doc: Document):
    return list(doc.element.body)


def _table_after_paragraph(doc: Document, paragraph: Paragraph) -> Optional[Table]:
    seen = False
    for child in _body_children(doc):
        if child is paragraph._p:
            seen = True
            continue
        if seen and child.tag == qn("w:tbl"):
            return Table(child, doc)
    return None


def _find_table(doc: Document, spec: Dict[str, Any]) -> Optional[Table]:
    columns = [str(c) for c in spec.get("columns") or []]
    anchor = spec.get("anchor")
    if anchor:
        for paragraph in doc.paragraphs:
            if str(anchor) in paragraph.text:
                found = _table_after_paragraph(doc, paragraph)
                if found is not None:
                    return found
    best = None
    best_score = -1
    for table in doc.tables:
        score = _match_score(_header_texts(table), columns)
        if score > best_score:
            best_score = score
            best = table
    return best


def _text_width_pt(doc: Document) -> float:
    section = doc.sections[0]
    try:
        return float(section.page_width.pt - section.left_margin.pt - section.right_margin.pt)
    except Exception:
        return TEXT_WIDTH_FALLBACK_PT


def _set_run_font(run, name: str, size_pt: float, bold: bool) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def _set_cell_text(cell, value: Any, font: str, size_pt: float, bold: bool,
                   line_pt: Optional[float]) -> None:
    text = "" if value is None else str(value)
    paragraph = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        parent = extra._element.getparent()
        if parent is not None:
            parent.remove(extra._element)
    for run in list(paragraph.runs):
        parent = run._element.getparent()
        if parent is not None:
            parent.remove(run._element)
    parts = text.split("\v")
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        _set_run_font(run, font, size_pt, bold)
        if index < len(parts) - 1:
            run.add_break()
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = 1
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    if line_pt:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line_pt)
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    v_align = tcpr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tcpr.append(v_align)
    v_align.set(qn("w:val"), "center")


def _grid_cols(table: Table):
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    return grid, grid.findall(qn("w:gridCol"))


def _unique_cells(row) -> List:
    seen = set()
    cells = []
    for cell in row.cells:
        cid = id(cell._tc)
        if cid in seen:
            continue
        seen.add(cid)
        cells.append(cell)
    return cells


def _add_column(table: Table) -> None:
    grid, cols = _grid_cols(table)
    if cols:
        grid.append(deepcopy(cols[-1]))
    else:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), "800")
        grid.append(col)
    for tr in table._tbl.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        if not tcs:
            continue
        new_tc = deepcopy(tcs[-1])
        for t in new_tc.findall(".//" + qn("w:t")):
            t.text = ""
        tr.append(new_tc)


def _delete_last_column(table: Table) -> None:
    grid, cols = _grid_cols(table)
    if cols:
        grid.remove(cols[-1])
    for tr in table._tbl.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        if tcs:
            tr.remove(tcs[-1])


def _set_widths(table: Table, widths_pt: Sequence[float]) -> None:
    total = sum(widths_pt)
    tbl = table._tbl
    tblpr = tbl.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        tbl.insert(0, tblpr)
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(total * 20)))
    tblw.set(qn("w:type"), "dxa")
    jc = tblpr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblpr.append(jc)
    jc.set(qn("w:val"), "center")
    _, cols = _grid_cols(table)
    for col, width in zip(cols, widths_pt):
        col.set(qn("w:w"), str(int(width * 20)))
    for tr in tbl.findall(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        for tc, width in zip(tcs, widths_pt):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr")
                tc.insert(0, tcpr)
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(int(width * 20)))
            tcw.set(qn("w:type"), "dxa")


def _mark_header_row(table: Table) -> None:
    tr = table.rows[0]._tr
    trpr = tr.find(qn("w:trPr"))
    if trpr is None:
        trpr = OxmlElement("w:trPr")
        tr.insert(0, trpr)
    if trpr.find(qn("w:tblHeader")) is None:
        trpr.append(OxmlElement("w:tblHeader"))


def _clear_data_rows(table: Table) -> None:
    trs = table._tbl.findall(qn("w:tr"))
    for tr in trs[1:]:
        table._tbl.remove(tr)


def _fill_one_table(doc: Document, table: Table, columns: List[str],
                    rows: List[List[Any]]) -> None:
    while len(_grid_cols(table)[1]) > len(columns):
        _delete_last_column(table)
    while len(_grid_cols(table)[1]) < len(columns):
        _add_column(table)
    is_detail = bool(columns) and columns[0].strip() == "序号"
    font = "宋体"
    size = DETAIL_FONT_PT if is_detail else _MATRIX_FONT_PT
    line = DETAIL_LINE_PT if is_detail else None
    header_cells = _unique_cells(table.rows[0])
    for index, name in enumerate(columns):
        if index < len(header_cells):
            _set_cell_text(header_cells[index], name, font, size, True, line)
    _clear_data_rows(table)
    for data in rows:
        row = table.add_row()
        cells = _unique_cells(row)
        for index, value in enumerate(data):
            if index < len(cells):
                _set_cell_text(cells[index], value, font, size, False, line)
    if is_detail:
        _mark_header_row(table)
        widths = detail_col_widths(columns, _text_width_pt(doc))
    else:
        text_w = _text_width_pt(doc)
        unit_n = max(1, len(columns) - 1)
        unit_w = (text_w - _MATRIX_FIRST_PT) / unit_n
        widths = [_MATRIX_FIRST_PT] + [unit_w] * unit_n
    _set_widths(table, widths)


def _insert_after(element, new_elm) -> None:
    element.addnext(new_elm)


def _caption_paragraph(text: str) -> Any:
    paragraph = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    ppr.append(jc)
    paragraph.append(ppr)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), "宋体")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "21")
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    paragraph.append(run)
    return paragraph


def _append_matrix_block(doc: Document, prev: Table, proto: Any,
                         columns: List[str], rows: List[List[Any]]) -> Table:
    mark = _caption_paragraph("（续表）")
    _insert_after(prev._tbl, mark)
    cloned = deepcopy(proto)
    _insert_after(mark, cloned)
    table = Table(cloned, doc)
    _fill_one_table(doc, table, columns, rows)
    return table


def _remove_table(table: Table) -> None:
    parent = table._tbl.getparent()
    if parent is not None:
        parent.remove(table._tbl)


def fill_tables(doc: Document, specs: Sequence[Dict[str, Any]]) -> None:
    for spec in specs:
        columns = [str(c) for c in spec.get("columns") or []]
        rows = spec.get("rows") or []
        if not columns:
            continue
        target = _find_table(doc, spec)
        if target is None:
            raise RuntimeError(f"未找到匹配附表：{columns}")
        if not rows:
            _remove_table(target)
            continue
        max_units = spec.get("max_units_per_table")
        if max_units is not None and len(columns) - 1 > int(max_units):
            proto = deepcopy(target._tbl)
            units = columns[1:]
            limit = int(max_units)
            count = -(-len(units) // limit)
            base, rem = divmod(len(units), count)
            blocks = []
            cursor = 0
            for index in range(count):
                size = base + (1 if index < rem else 0)
                blocks.append(units[cursor:cursor + size])
                cursor += size
            prev = target
            start = 0
            for bi, block in enumerate(blocks):
                blk_cols = [columns[0]] + block
                blk_rows = [[row[0]] + row[1 + start:1 + start + len(block)] for row in rows]
                if bi == 0:
                    _fill_one_table(doc, target, blk_cols, blk_rows)
                    prev = target
                else:
                    prev = _append_matrix_block(doc, prev, proto, blk_cols, blk_rows)
                start += len(block)
        else:
            _fill_one_table(doc, target, columns, rows)


def trim_month_weeks(doc: Document, keep_weeks: int) -> None:
    if keep_weeks >= 5:
        return
    paragraphs = list(doc.paragraphs)
    start = end = None
    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        matched = _WEEK_RE.match(text)
        if matched and int(matched.group(2)) > keep_weeks and start is None:
            start = index
        if text.startswith("三、") and start is not None:
            end = index
            break
    if start is None or end is None or end <= start:
        return
    for paragraph in paragraphs[start:end]:
        parent = paragraph._element.getparent()
        if parent is not None:
            parent.remove(paragraph._element)


def _png_size(path: str) -> Tuple[float, float]:
    from PIL import Image
    with Image.open(path) as image:
        return float(image.size[0]), float(image.size[1])


def _format_figure_paragraph(paragraph: Paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.append(jc)
    jc.set(qn("w:val"), "center")
    for tag in ("w:keepNext", "w:keepLines"):
        el = ppr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            ppr.append(el)


def _clear_drawings(paragraph: Paragraph) -> None:
    for drawing in list(paragraph._p.findall(".//" + qn("w:drawing"))):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)


def put_picture(paragraph: Paragraph, png: str) -> None:
    if not os.path.isfile(png):
        LOGGER.warning("图表不存在，跳过：%s", png)
        return
    _clear_drawings(paragraph)
    _format_figure_paragraph(paragraph)
    width, height = _CHART_W_PT, _CHART_W_PT * 0.42
    try:
        pw, ph = _png_size(png)
        if pw:
            height = _CHART_W_PT * (ph / pw)
            if height > _MAX_H_PT:
                height = _MAX_H_PT
                width = _MAX_H_PT * (pw / ph)
    except Exception as exc:
        LOGGER.warning("读取 PNG 尺寸失败：%s", exc)
    run = paragraph.add_run()
    run.add_picture(png, width=Pt(width), height=Pt(height))


def _insert_picture_after(paragraph: Paragraph, png: str) -> None:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    put_picture(Paragraph(new_p, paragraph._parent), png)


def choose_cap_png(caption: Optional[str], index: int, report: str, cap_dir: str) -> Optional[str]:
    if report != "month":
        if index == 1:
            prefix = _SINGLE_PREFIX.get(report, "cap_w1")
            return os.path.join(cap_dir, f"{prefix}_单位承载力.png")
        return None
    if not caption:
        return None
    if "作业类型" in caption:
        return os.path.join(cap_dir, "cap_m_作业类型.png")
    week = next((code for key, code in _WEEK_PREFIX if key in caption), None)
    head = "管理承载力" if "管理承载力" in caption else "单位承载力"
    return os.path.join(cap_dir, f"cap_{week if week else 'm'}_{head}.png")


def _first_drawing_paragraph(doc: Document) -> Optional[Paragraph]:
    for paragraph in doc.paragraphs:
        if paragraph._p.find(".//" + qn("w:drawing")) is not None:
            return paragraph
    return None


def apply_charts(doc: Document, report: str, cap_dir: str) -> int:
    """按 pngCharts.apply 的图位规则插入 PNG。返回插入张数。"""
    count = 0
    if report == "month":
        paragraphs = list(doc.paragraphs)
        for index, paragraph in enumerate(paragraphs):
            text = paragraph.text.strip()
            if text.startswith("图 ") and len(text) <= 40 and index > 0:
                png = choose_cap_png(text, index, report, cap_dir)
                if png:
                    put_picture(paragraphs[index - 1], png)
                    count += 1
        return count
    slot = _first_drawing_paragraph(doc)
    png = choose_cap_png(None, 1, report, cap_dir)
    if slot is not None and png:
        put_picture(slot, png)
        count += 1
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip().startswith(_PTYPE_ANCHOR):
            _insert_picture_after(
                paragraph, os.path.join(cap_dir, f"{_SINGLE_PREFIX.get(report, 'cap_w1')}_作业类型.png"))
            count += 1
            break
    keyword = _MANAGE_ANCHOR.get(report)
    if keyword:
        for paragraph in list(doc.paragraphs):
            if keyword in paragraph.text:
                _insert_picture_after(
                    paragraph,
                    os.path.join(cap_dir, f"{_SINGLE_PREFIX.get(report, 'cap_w1')}_管理承载力.png"))
                count += 1
                break
    return count


def render(template_docx: str, params: Dict[str, Any], out_path: str,
           report: str, charts_dir: Optional[str]) -> int:
    doc = Document(template_docx)
    apply_replace(doc, params.get("replace") or {})
    apply_paragraph_replace(doc, params.get("paragraph_replace") or [])
    repair_broken_refs(doc)
    fill_tables(doc, params.get("tables") or [])
    keep = params.get("trim_month_weeks")
    if keep is not None:
        trim_month_weeks(doc, int(keep))
    charts = 0
    if charts_dir:
        charts = apply_charts(doc, report, charts_dir)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    LOGGER.info("已生成 %s（图 %d 张）", out_path, charts)
    return charts


def convert_doc_to_docx(doc_path: str, out_dir: str) -> str:
    """用 LibreOffice 把 .doc 模板转成 .docx，不改原模板。"""
    os.makedirs(out_dir, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "docx",
         "--outdir", out_dir, doc_path],
        check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    name = os.path.splitext(os.path.basename(doc_path))[0] + ".docx"
    produced = os.path.join(out_dir, name)
    if not os.path.isfile(produced):
        raise FileNotFoundError(produced)
    return produced


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="fillReportLinux.py")
    parser.add_argument("--template", required=True, choices=list(TEMPLATES))
    parser.add_argument("--params", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--report", required=True, choices=("day", "week", "month"))
    parser.add_argument("--charts", default=None)
    parser.add_argument("--docx-template", default=None,
                        help="已转换的 docx 模板；缺省时从 templates/*.doc 用 soffice 转换")
    return parser


def main(argv: Optional[List[str]] = None) -> int:
    logging.basicConfig(level=logging.INFO, stream=__import__("sys").stderr,
                        format="[%(levelname)s] %(message)s")
    args = build_parser().parse_args(argv)
    import json
    with open(args.params, "r", encoding="utf-8") as handle:
        params = json.load(handle)
    docx_template = args.docx_template
    if not docx_template:
        skill_templates = os.path.abspath(os.path.join(
            os.path.dirname(__file__), "..", "templates"))
        doc_name = TEMPLATES[args.template]
        doc_path = os.path.join(skill_templates, doc_name)
        work = tempfile.mkdtemp(prefix="cap-docx-")
        try:
            docx_template = convert_doc_to_docx(doc_path, work)
            render(docx_template, params, args.out, args.report, args.charts)
        finally:
            shutil.rmtree(work, ignore_errors=True)
    else:
        render(docx_template, params, args.out, args.report, args.charts)
    print(json.dumps({"ok": True, "out": os.path.abspath(args.out)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
